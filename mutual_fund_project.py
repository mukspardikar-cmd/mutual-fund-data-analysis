"""
Mutual Fund Overview & Insights  -  end-to-end analysis script
Run:  python mutual_fund_project.py
Keep 'comprehensive_mutual_funds_data.csv' in the SAME folder as this file.

Outputs created in the same folder:
  cleaned_funds.csv      -> load this into Power BI
  top30_funds.xlsx       -> Stage 5 deliverable
  charts/*.png           -> Stage 2 plots
  (all key numbers for the write-up are printed in the terminal)
"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")              # saves charts to files (no pop-up windows)
import matplotlib.pyplot as plt
from sklearn.preprocessing import MinMaxScaler

os.makedirs("charts", exist_ok=True)
pd.set_option("display.width", 200)
pd.set_option("display.max_columns", 50)

# ----------------------------------------------------------------------
# STAGE 1 : DATA CLEANING
# ----------------------------------------------------------------------
df = pd.read_csv("comprehensive_mutual_funds_data.csv")
print("\n=== STAGE 1 ===")
print("Rows, Columns:", df.shape)
print(df.dtypes)

# 1.3  Missing values. Some columns hide missing data as the text "-"
RATIO_COLS = ["sortino", "alpha", "sd", "beta", "sharpe"]
for c in RATIO_COLS:
    df[c] = pd.to_numeric(df[c].replace("-", np.nan), errors="coerce")
print("\nMissing values per column:")
print(df.isna().sum()[lambda s: s > 0])
print("\n% missing:")
print((df.isna().mean() * 100).round(1)[lambda s: s > 0])

# 1.2  Drop columns we do not need (risk is already captured by risk_level)
df = df.drop(columns=RATIO_COLS + ["returns_5yr"])
# returns_5yr dropped: ~20% missing and the project scores on 3-year returns.

# Strategy: funds without 3-year return are too new to score -> drop them
before = len(df)
df = df.dropna(subset=["returns_3yr"]).copy()
print(f"\nDropped {before - len(df)} funds with no 3-year return")

# Data-quality check (found while exploring): the value 32.4 shows up as the
# 3-year return for 10 unrelated funds (Debt, Hybrid, Equity) and a 2-year-old
# fund cannot have a 3-year return at all. These look like placeholder/bad data.
# Set EXCLUDE_SUSPECT = False if you want to keep them and mention it in your write-up.
EXCLUDE_SUSPECT = True
suspect = (df["returns_3yr"] == 32.4) | (df["fund_age_yr"] < 3)
print(f"Suspect rows (3yr return = 32.4 or age < 3): {suspect.sum()}")
if EXCLUDE_SUSPECT:
    df = df[~suspect].copy()

# 1.4  Numeric formats: clean text, make sure numbers are numbers
df["scheme_name"] = df["scheme_name"].str.strip()
for c in ["expense_ratio", "returns_1yr", "returns_3yr", "fund_size_cr",
          "min_sip", "min_lumpsum", "fund_age_yr", "risk_level", "rating"]:
    df[c] = pd.to_numeric(df[c], errors="coerce")

# Friendly labels (risk_level 1 = lowest risk, 6 = highest risk)
risk_map = {1: "Low", 2: "Low to Moderate", 3: "Moderate",
            4: "Moderately High", 5: "High", 6: "Very High"}
df["risk_label"] = df["risk_level"].map(risk_map)
# the dataset has no 'fund type' column; 'category' (Equity/Debt/Hybrid...) is used as fund type
df["fund_type"] = df["category"]

# ----------------------------------------------------------------------
# STAGE 2 : DATA DESCRIPTION
# ----------------------------------------------------------------------
print("\n=== STAGE 2 ===")
num_cols = ["returns_1yr", "returns_3yr", "expense_ratio",
            "fund_age_yr", "risk_level", "fund_size_cr"]
summary = df[num_cols].agg(["mean", "median", "min", "max", "std"]).T
summary["mode"] = [df[c].mode().iloc[0] for c in num_cols]
print(summary.round(2))
print("Skewness:\n", df[num_cols].skew().round(2))

for col in ["returns_3yr", "expense_ratio", "fund_age_yr", "risk_level"]:
    plt.figure(figsize=(6, 4))
    df[col].plot(kind="hist", bins=20, edgecolor="black")
    plt.title(f"Distribution of {col}")
    plt.xlabel(col)
    plt.tight_layout()
    plt.savefig(f"charts/hist_{col}.png", dpi=120)
    plt.close()

plt.figure(figsize=(6, 4))
plt.boxplot(df["returns_1yr"], vert=False)
plt.title("1-year return - outliers")
plt.tight_layout()
plt.savefig("charts/box_returns_1yr.png", dpi=120)
plt.close()

# Outliers using the IQR rule
q1, q3 = df["returns_1yr"].quantile([0.25, 0.75])
iqr = q3 - q1
out = df[(df["returns_1yr"] > q3 + 1.5 * iqr) | (df["returns_1yr"] < q1 - 1.5 * iqr)]
print(f"\n1-yr return outliers: {len(out)}")
print(out.sort_values("returns_1yr", ascending=False)[
    ["scheme_name", "category", "returns_1yr"]].head(5).to_string(index=False))
print("\nFund age counts (note the pile-up at 10):")
print(df["fund_age_yr"].value_counts().sort_index().to_string())

# ----------------------------------------------------------------------
# STAGE 3 : NORMALISATION (MinMaxScaler -> 0 to 1)
# ----------------------------------------------------------------------
print("\n=== STAGE 3 ===")
# 1-yr return has an extreme outlier (130%). Clip to 1st-99th percentile
# ONLY for scaling so one weird fund does not squash everyone else.
lo, hi = df["returns_1yr"].quantile([0.01, 0.99])
df["returns_1yr_clipped"] = df["returns_1yr"].clip(lo, hi)

# Moderate age: closest to the 'ideal' age scores highest (triangle shape)
IDEAL_AGE = 6
df["age_distance"] = (df["fund_age_yr"] - IDEAL_AGE).abs()

to_scale = ["returns_3yr", "returns_1yr_clipped", "expense_ratio",
            "age_distance", "risk_level"]
scaler = MinMaxScaler()
scaled = pd.DataFrame(scaler.fit_transform(df[to_scale]),
                      columns=[c + "_n" for c in to_scale], index=df.index)
df = pd.concat([df, scaled], axis=1)
print(df[[c + "_n" for c in to_scale]].describe().loc[["min", "max", "mean"]].round(2))

# ----------------------------------------------------------------------
# STAGE 4 : SCORING & RANKING
# ----------------------------------------------------------------------
print("\n=== STAGE 4 ===")
# Direction handling: "higher is better" columns are used as they are,
# "lower is better" columns are flipped with (1 - x).
df["s_return3"] = df["returns_3yr_n"]                  # higher = better
df["s_expense"] = 1 - df["expense_ratio_n"]            # lower  = better
df["s_age"] = 1 - df["age_distance_n"]                 # closer to ideal = better
df["s_consistent"] = (df["returns_1yr"] > 0).astype(int) * (
    0.5 + 0.5 * df["returns_1yr_clipped_n"])           # 0 if negative, else 0.5-1.0
df["s_risk"] = 1 - df["risk_level_n"]                  # lower risk = better

WEIGHTS = {"s_return3": 0.40, "s_expense": 0.20, "s_age": 0.10,
           "s_consistent": 0.15, "s_risk": 0.15}


def score(frame, w):
    return sum(frame[k] * v for k, v in w.items()) * 100


df["score"] = score(df, WEIGHTS)
df = df.sort_values("score", ascending=False).reset_index(drop=True)
df["rank"] = df.index + 1
print("Weights:", WEIGHTS)

# Which features drive the score? (correlation of each part with final score)
print("\nCorrelation of each factor with final score:")
print(df[list(WEIGHTS)].corrwith(df["score"]).round(2).sort_values(ascending=False))

# Sensitivity: does a different weighting change the Top 30?
alt = {"s_return3": 0.25, "s_expense": 0.20, "s_age": 0.10,
       "s_consistent": 0.15, "s_risk": 0.30}   # risk-averse investor
alt_top = set(df.assign(sc=score(df, alt)).nlargest(30, "sc")["scheme_name"])
overlap = len(alt_top & set(df.head(30)["scheme_name"]))
print(f"\nSensitivity: risk-heavy weighting keeps {overlap}/30 of the same funds")

# ----------------------------------------------------------------------
# STAGE 5 : TOP 30 -> EXCEL
# ----------------------------------------------------------------------
print("\n=== STAGE 5 ===")
top30 = df.head(30)
cols = ["rank", "scheme_name", "amc_name", "category", "sub_category",
        "risk_label", "rating", "fund_age_yr", "expense_ratio",
        "returns_1yr", "returns_3yr", "fund_size_cr", "min_sip", "score"]
top_out = top30[cols].copy()
top_out["score"] = top_out["score"].round(2)

with pd.ExcelWriter("top30_funds.xlsx", engine="openpyxl") as xw:
    top_out.to_excel(xw, sheet_name="Top 30", index=False)
    ws = xw.sheets["Top 30"]
    from openpyxl.styles import Font, PatternFill, Alignment
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 3, 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

print(top30[["rank", "scheme_name", "category", "risk_label",
             "expense_ratio", "returns_3yr", "score"]].to_string(index=False))
print("\nTop-30 by category:\n", top30["category"].value_counts().to_string())
print("\nTop-30 by risk:\n", top30["risk_label"].value_counts().to_string())
print("\nTop-30 by AMC:\n", top30["amc_name"].value_counts().head(5).to_string())

# ----------------------------------------------------------------------
# Export cleaned data for Power BI
# ----------------------------------------------------------------------
keep = ["scheme_name", "fund_type", "category", "sub_category", "amc_name",
        "fund_manager", "risk_level", "risk_label", "rating", "fund_age_yr",
        "expense_ratio", "fund_size_cr", "min_sip", "min_lumpsum",
        "returns_1yr", "returns_3yr", "score", "rank"]
df[keep].to_csv("cleaned_funds.csv", index=False)

# ----------------------------------------------------------------------
# STAGE 7 : NUMBERS FOR THE WRITE-UP
# ----------------------------------------------------------------------
print("\n=== STAGE 7 ANSWERS ===")
aum_type = df.groupby("fund_type")["fund_size_cr"].sum().sort_values(ascending=False)
print("1) AUM by fund type (Rs Cr):\n", aum_type.round(0).to_string())
print(f"   Leader: {aum_type.index[0]}, leads #2 by {aum_type.iloc[0] - aum_type.iloc[1]:,.0f} Cr")

mgr = df.groupby("fund_manager")["fund_size_cr"].sum().sort_values(ascending=False)
print("\n2) Top fund managers by AUM (Rs Cr):\n", mgr.head(3).round(0).to_string())

er_cat = df.groupby("fund_type")["expense_ratio"].mean().sort_values()
er_sub = df.groupby("sub_category")["expense_ratio"].mean().sort_values()
print("\n3) Avg expense ratio by fund type:\n", er_cat.round(2).to_string())
print("   Lowest sub-category:", er_sub.index[0], round(er_sub.iloc[0], 2))

best1 = df.loc[df["returns_1yr"].idxmax()]
print("\n4) Best 1-yr return:", best1["scheme_name"], best1["returns_1yr"], "%")
valid = df[df["scheme_name"] != best1["scheme_name"]].nlargest(1, "returns_1yr").iloc[0]
print("   (Outlier - next best:", valid["scheme_name"], valid["returns_1yr"], "%)")

sip = df.loc[df["min_sip"] > 0, "min_sip"]
print(f"\n5) Avg min SIP (SIP funds only): {sip.mean():.0f}  | median {sip.median():.0f}")
print(f"   Median min lump-sum: {df['min_lumpsum'].median():.0f} | mean {df['min_lumpsum'].mean():.0f}")

print("\n6) Return vs risk by category:")
print(df.groupby("fund_type").agg(avg_3yr=("returns_3yr", "mean"),
                                  avg_risk_level=("risk_level", "mean"),
                                  avg_expense=("expense_ratio", "mean"),
                                  funds=("scheme_name", "count")).round(2).to_string())
print("\nDone. Files: cleaned_funds.csv, top30_funds.xlsx, charts/")
